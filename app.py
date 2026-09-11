import csv
import os
import re
from datetime import datetime

from flask import Flask, jsonify, render_template, request
from openpyxl import load_workbook

import gservices
import gsheet_db as db

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_PATH = os.path.join(BASE_DIR, "MANPOWER HC.xlsx")
UPLOAD_DIR = os.path.join(BASE_DIR, "uploads")

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 32 * 1024 * 1024
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "manpower-sheets")

os.makedirs(UPLOAD_DIR, exist_ok=True)


# ------------------------------------------------------------------
# Biometric log parsing (flexible auto-detect)
# ------------------------------------------------------------------
def looks_like_key(s, keys):
    s = str(s).strip().lower().replace("_", "").replace("-", "").replace(" ", "")
    return s in keys or any(k in s for k in keys)


def split_datetime(s):
    s = str(s).strip()
    for fmt in ("%m/%d/%Y %H:%M", "%m/%d/%Y %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S",
                "%m-%d-%Y %H:%M", "%m-%d-%Y %H:%M:%S", "%Y/%m/%d %H:%M", "%Y/%m/%d %H:%M:%S",
                "%d/%m/%Y %H:%M", "%d/%m/%Y %H:%M:%S", "%d-%m-%Y %H:%M", "%d-%m-%Y %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%m-%d-%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.combine(datetime.strptime(s, fmt).date(), datetime.min.time())
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(s)
    except ValueError:
        return None


def parse_date_and_time(dval, tval):
    if dval is None:
        return None
    d = str(dval).strip()
    if tval is not None:
        t = str(tval).strip()
        for dfmt in ("%m/%d/%Y", "%Y-%m-%d", "%m-%d-%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
            try:
                dd = datetime.strptime(d, dfmt)
            except ValueError:
                continue
            for tfmt in ("%H:%M:%S", "%H:%M"):
                try:
                    tt = datetime.strptime(t, tfmt)
                    return dd.replace(hour=tt.hour, minute=tt.minute, second=tt.second)
                except ValueError:
                    continue
        # date part of a datetime string
        dt = split_datetime(d)
        if dt is not None:
            try:
                ttt = datetime.strptime(t, "%H:%M:%S")
            except ValueError:
                try:
                    ttt = datetime.strptime(t, "%H:%M")
                except ValueError:
                    return dt
            return dt.replace(hour=ttt.hour, minute=ttt.minute, second=ttt.second)
        return None
    return split_datetime(dval)


def parse_biometric_file(path, filename=""):
    """Try to parse a biometric log file. Returns list of raw records."""
    ext = os.path.splitext(filename)[1].lower()
    raw_rows = []

    if ext in (".xlsx", ".xls"):
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        for row in ws.iter_rows(values_only=True):
            if row and any(c is not None for c in row):
                raw_rows.append(["" if c is None else str(c).strip() for c in row])
    elif ext == ".csv":
        with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
            sample = f.read(4096)
        if "," in sample:
            delim = ","
        elif ";" in sample:
            delim = ";"
        elif "\t" in sample:
            delim = "\t"
        else:
            delim = ","
        with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
            reader = csv.reader(f, delimiter=delim)
            for row in reader:
                if row and any(c.strip() for c in row):
                    raw_rows.append([c.strip() for c in row])
    else:
        with open(path, "r", encoding="utf-8-sig", errors="replace") as f:
            for line in f:
                if not line.strip():
                    continue
                if "," in line:
                    parts = [p.strip() for p in line.split(",")]
                elif "\t" in line:
                    parts = [p.strip() for p in line.split("\t")]
                else:
                    parts = [p.strip() for p in line.split()]
                raw_rows.append(parts)

    return normalize_rows(raw_rows)


def normalize_rows(raw_rows):
    """Auto-detect columns from raw parsed rows."""
    if not raw_rows:
        return []
    header = None
    start_idx = 0
    for i, row in enumerate(raw_rows[:20]):
        joined = " ".join(row).lower()
        if any(k in joined for k in ("date", "time", "user", "id", "name", "punch", "check", "record")):
            header = row
            start_idx = i + 1
            break

    if header is None:
        header = ["id", "datetime"]
        start_idx = 0

    h = [str(x).strip().lower() for x in header]

    def find_col(keys):
        for idx, name in enumerate(h):
            if isinstance(name, str) and looks_like_key(name, keys):
                return idx
        return None

    i_dt = find_col(["datetime", "dateandtime", "date and time", "attendance time", "punchtime", "recordtime", "record time"])
    i_date = find_col(["date"])
    i_time = find_col(["time"])
    i_id = find_col(["userid", "user id", "userno", "eno", "employeeid", "employee id", "empid", "badge", "idno", "id no", "enrollid", "srno", "id"])
    i_name = find_col(["name", "employeename", "employee name", "fullname"])
    i_fname = find_col(["firstname", "first name", "fname", "givenname"])
    i_lname = find_col(["lastname", "last name", "lname", "surname"])
    i_type = find_col(["verified", "verification", "recordstate", "record state", "state", "punchtype", "type", "status", "inout", "in/out", "io"])

    # ZKTeco ADMS .dat files have no header. Detect the in/out column by finding an
    # all-integer column (values in 0..5) that varies, after id/date/time columns.
    if i_type is None:
        col_values = {}
        for row in raw_rows[start_idx:start_idx + 200]:
            for ci, val in enumerate(row):
                col_values.setdefault(ci, set()).add(str(val).strip())
        for ci, vals in col_values.items():
            if ci in (i_id, i_dt, i_date, i_time):
                continue
            if vals and vals <= {"0", "1", "2", "3", "4", "5"} and len(vals) >= 2:
                i_type = ci
                break

    records = []
    for row in raw_rows[start_idx:]:
        if not row or all(not str(c).strip() for c in row):
            continue
        pad = row + [""] * (len(h) - len(row))

        dt = None
        if i_dt is not None:
            dt = split_datetime(pad[i_dt])
        elif i_date is not None:
            dt = parse_date_and_time(pad[i_date], pad[i_time] if i_time is not None else None)

        if dt is None:
            continue

        emp_key = ""
        if i_id is not None:
            emp_key = str(pad[i_id]).strip()
        # some files put the id as first cell when no header detected
        if not emp_key and header[0] in ("id", "emp"):
            emp_key = str(row[0]).strip()

        name = ""
        if i_name is not None:
            name = str(pad[i_name]).strip()
        elif i_fname is not None or i_lname is not None:
            fn = str(pad[i_fname]).strip() if i_fname is not None else ""
            ln = str(pad[i_lname]).strip() if i_lname is not None else ""
            name = (fn + " " + ln).strip()

        ptype = ""
        if i_type is not None:
            t = str(pad[i_type]).strip()
            tl = t.lower()
            # ZKTeco in/out state codes: 0 check-in, 1 check-out, 2 break-out,
            # 3 break-in, 4 overtime-in, 5 overtime-out
            zk_state = {"0": "IN", "1": "OUT", "2": "OUT", "3": "IN", "4": "IN", "5": "OUT"}
            if tl in zk_state:
                ptype = zk_state[tl]
            elif tl in ("in", "checkin", "check in", "entry", "i"):
                ptype = "IN"
            elif tl in ("out", "checkout", "check out", "exit", "o"):
                ptype = "OUT"
            else:
                ptype = t

        records.append({
            "emp_key": emp_key,
            "name": name,
            "punch_time": dt.strftime("%Y-%m-%d %H:%M:%S"),
            "punch_type": ptype,
        })
    return records


# ------------------------------------------------------------------
# Routes: UI
# ------------------------------------------------------------------
@app.route("/")
def index():
    return render_template("index.html")


# ------------------------------------------------------------------
# Routes: Workers API
# ------------------------------------------------------------------
@app.route("/api/workers", methods=["GET"])
def list_workers():
    return jsonify(db.get_workers())


@app.route("/api/workers", methods=["POST"])
def add_worker():
    data = request.get_json(force=True)
    lname = str(data.get("lname", "")).strip()
    fname = str(data.get("fname", "")).strip()
    if not lname and not fname:
        return jsonify({"ok": False, "error": "Name is required"}), 400
    try:
        wid = db.add_worker(
            str(data.get("employee_id", "")).strip(),
            lname, fname,
            str(data.get("position", "")).strip(),
        )
        return jsonify({"ok": True, "id": wid})
    except ValueError as e:
        return jsonify({"ok": False, "error": str(e)}), 400


@app.route("/api/workers/<int:wid>", methods=["PUT"])
def update_worker(wid):
    data = request.get_json(force=True)
    try:
        db.update_worker(wid, data)
        return jsonify({"ok": True})
    except KeyError:
        return jsonify({"ok": False, "error": "Not found"}), 404
    except ValueError as e:
        return jsonify({"ok": False, "error": str(e)}), 400


@app.route("/api/workers/<int:wid>", methods=["DELETE"])
def delete_worker(wid):
    db.delete_worker(wid)
    return jsonify({"ok": True})


@app.route("/api/workers/batch-employee-id", methods=["POST"])
def batch_employee_id():
    data = request.get_json(force=True)
    pairs = data.get("pairs", [])
    if not isinstance(pairs, list) or not pairs:
        return jsonify({"ok": False, "error": "No data"}), 400
    updated = db.batch_set_employee_ids(pairs)
    return jsonify({"ok": True, "updated": updated})


# ------------------------------------------------------------------
# Routes: Seeding
# ------------------------------------------------------------------
@app.route("/api/seed", methods=["POST"])
def seed():
    if not os.path.exists(EXCEL_PATH):
        return jsonify({"ok": False, "error": "Excel file not found: " + EXCEL_PATH}), 400
    wb = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    ws = wb.active
    headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
    col_idx = {h: i for i, h in enumerate(headers)}

    def col(name):
        for cname, idx in col_idx.items():
            if cname == name or name in cname:
                return idx
        return None

    i_id, i_ln, i_fn, i_pos = col("employee_id"), col("lname"), col("fname"), col("position")
    seed_rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None or all(c is None for c in row):
            continue
        seed_rows.append((
            str(row[i_id] or "").strip() if i_id is not None else "",
            str(row[i_ln] or "").strip() if i_ln is not None else "",
            str(row[i_fn] or "").strip() if i_fn is not None else "",
            str(row[i_pos] or "").strip() if i_pos is not None else "",
        ))
    return jsonify(db.seed_workers(seed_rows))


@app.route("/api/reset", methods=["POST"])
def reset():
    db.reset_all()
    return jsonify({"ok": True})


@app.route("/api/stats", methods=["GET"])
def stats():
    return jsonify(db.stats())


# ------------------------------------------------------------------
# Routes: Attendance import
# ------------------------------------------------------------------
@app.route("/api/attendance/import", methods=["POST"])
def import_attendance():
    if "file" not in request.files:
        return jsonify({"ok": False, "error": "No file uploaded"}), 400
    f = request.files["file"]
    if f.filename == "":
        return jsonify({"ok": False, "error": "No file selected"}), 400
    save_path = os.path.join(UPLOAD_DIR, datetime.now().strftime("%Y%m%d%H%M%S") + "_" + f.filename)
    f.save(save_path)
    records = parse_biometric_file(save_path, f.filename)
    if not records:
        return jsonify({"ok": False, "error": "Could not parse any records from this file. Check the format."}), 400

    # Keep a copy of the uploaded log in Google Drive.
    try:
        gservices.upload_file_to_drive(save_path, f.filename)
    except Exception as e:
        app.logger.warning("Drive upload failed: %s", e)

    matched, unmatched = [], []
    for rec in records:
        w = db.find_worker(rec["emp_key"]) or db.worker_by_name(rec["name"])
        if w:
            matched.append({**rec, "worker_id": w["id"], "employee_id": w["employee_id"],
                            "name": f"{w['lname']}, {w['fname']}"})
        else:
            unmatched.append(rec)
    return jsonify({
        "ok": True,
        "file": f.filename,
        "total_records": len(records),
        "matched": matched,
        "unmatched": unmatched,
        "preview": matched[:200],
    })


@app.route("/api/attendance/import/commit", methods=["POST"])
def commit_import():
    data = request.get_json(force=True)
    records = data.get("records", [])
    filename = data.get("filename", "")
    mapping = data.get("mapping", {})
    if not records:
        return jsonify({"ok": False, "error": "No records"}), 400
    result = db.commit_attendance(records, filename, mapping)
    if result["ok"]:
        try:
            backup = gservices.backup_spreadsheet()
            result["backup"] = backup
        except Exception as e:
            app.logger.warning("Spreadsheet backup failed: %s", e)
    return jsonify(result)


# ------------------------------------------------------------------
# Routes: Reports
# ------------------------------------------------------------------
@app.route("/api/reports/daily", methods=["GET"])
def daily_report():
    date = request.args.get("date", datetime.now().date().isoformat())
    return jsonify(db.daily_report(date))


@app.route("/api/reports/positions", methods=["GET"])
def positions_report():
    date = request.args.get("date", datetime.now().date().isoformat())
    return jsonify(db.positions_report(date))


# ------------------------------------------------------------------
# Routes: Google integration
# ------------------------------------------------------------------
@app.route("/api/google/status", methods=["GET"])
def google_status():
    try:
        gservices.get_credentials()
        folders = gservices.get_folders()
        return jsonify({
            "ok": True,
            "connected": True,
            "spreadsheet_url": gservices.spreadsheet_url(),
            "drive_root_url": "https://drive.google.com/drive/folders/" + folders["root_folder_id"],
            "attendance_folder_id": folders["attendance_folder_id"],
            "backup_folder_id": folders["backup_folder_id"],
        })
    except Exception as e:
        return jsonify({"ok": False, "connected": False, "error": str(e)}), 500


@app.route("/api/backup", methods=["POST"])
def backup():
    try:
        result = gservices.backup_spreadsheet()
        return jsonify({
            "ok": True,
            "name": result.get("name"),
            "url": result.get("webViewLink"),
            "file_id": result.get("id"),
        })
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500


# ------------------------------------------------------------------
if __name__ == "__main__":
    db.ensure_setup()
    print("=" * 60)
    print("MANPOWER SYSTEM (Google Sheets database)")
    print("Spreadsheet: " + gservices.SPREADSHEET_TITLE)
    print("Open:    http://127.0.0.1:5000")
    print("=" * 60)
    app.run(host="127.0.0.1", port=5000, debug=False)